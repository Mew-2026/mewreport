"""Weight reconciliation: why does the invoice bill more (or fewer) kg than were loaded?

  Loaded weight = sum over models of  loading-sheet qty x master kg/pc
  Billed weight = sum over invoice lines of  qty x invoice kg/pc,  invoice kg/pc = rate / Rate per kg to 2 decimals
  Variance      = billed - loaded, split exactly into
    Check 1  unit price / weight :  invoice qty x (invoice kg/pc - master kg/pc)   per model
    Check 2  quantity            :  (invoice qty - sheet qty) x master kg/pc        per model
    + billed kg of models with no weight at all (not in master, and nothing billed to take a weight from)

A model missing from the master is weighed with its invoice kg/pc (rate / Rate per kg), so its loaded pcs
still count; when it was not billed, the invoice kg/pc of its manual-match partner or of the closest billed
model (same Pro / non-Pro) is used. Its unit price cannot be checked, so Check 1 is 0 for it.

Works on the Step 1 rows (reconcile.reconcile, one row per model) and the parsed invoice lines, so
nothing is read again.
"""
from collections import Counter, defaultdict

from rapidfuzz import fuzz

from invoice_parser import model_sort_key
from mastersheet import code_key
from reconcile import SUGGEST_SCORE


def _model_kgs(master, model):
    """kg/pc of master rows whose description gives this model (most common first)."""
    return [k for k, _ in Counter(m["kg"] for m in master["rows"] if m["model"] == model).most_common()]


def _similar(key, candidates):
    """Closest billed model to weigh an unbilled model that is not in the master; Pro only with Pro."""
    pro = key.endswith("P")
    best = max(((fuzz.ratio(key, c), c) for c in candidates if c != key and c.endswith("P") == pro), default=(0, None))
    return best[1] if best[0] >= SUGGEST_SCORE else None


def weight_recon(recon_rows, inv, master, rate_per_kg):
    """One model can sit under several item codes with different weights (and rates), so Check 1 works
    line by line with each line's own item-code weight; the model's kg/pc for loaded weight and Check 2 is
    the average of the billed lines' weights (weighted by qty), or the master's weight when not billed."""
    lines_by_model = defaultdict(list)
    for l in inv["lines"]:
        lines_by_model[l["model"]].append(l)
    # invoice kg/pc per billed model (qty-weighted, each line's rate / Rate per kg to 2 dp)
    inv_kg = {}
    for mod, ls in lines_by_model.items():
        priced = [l for l in ls if l.get("rate")]
        if priced:
            inv_kg[mod] = (sum(l["qty"] * round(l["rate"] / rate_per_kg, 2) for l in priced)
                           / sum(l["qty"] for l in priced))
    group_models = defaultdict(list)   # manual match group -> invoice models in it
    for r in recon_rows:
        if r.get("group") and r["invoice_model"]:
            group_models[r["group"]].append(r["invoice_model"])

    models, not_in_master = [], []
    for r in recon_rows:
        key = r.get("model") or r["invoice_model"] or r["sheet_model"]
        lines = lines_by_model.get(r["invoice_model"], []) if r["invoice_model"] else []
        fallback = _model_kgs(master, key)
        line_kg = []   # (line, master kg/pc of its item code)
        for l in lines:
            m_ = master["by_code"].get(code_key(l["code"]))
            line_kg.append((l, m_["kg"] if m_ else (fallback[0] if fallback else None)))
        known_lines = [(l, k) for l, k in line_kg if k is not None]
        sq, iq = r["sheet_qty"], r["invoice_qty"]
        source = "master"
        if known_lines:
            kg = sum(l["qty"] * k for l, k in known_lines) / sum(l["qty"] for l, _ in known_lines)
        elif fallback:
            kg = fallback[0]
        else:   # not in the master: weigh with the invoice kg/pc instead of leaving the pcs out
            kg, source = None, None
            if r["invoice_model"] in inv_kg:
                kg, source = inv_kg[r["invoice_model"]], "invoice"
            else:
                partner = next((g for g in group_models.get(r.get("group"), []) if g in inv_kg), None) \
                    or _similar(key, inv_kg)
                if partner:
                    kg, source = inv_kg[partner], f"invoice {partner}"
        kgs = sorted({round(k, 4) for _, k in known_lines} | set(fallback if not known_lines else []))
        priced = [(l, k) for l, k in line_kg if l.get("rate")]
        billed = sum(l["qty"] * round(l["rate"] / rate_per_kg, 2) for l, _ in priced)   # kg/pc to 2 dp, like the master
        pq = sum(l["qty"] for l, _ in priced)
        m = {"model": key, "sheet_model": r["sheet_model"], "sheet_lines": r.get("sheet_lines", ""),
             "sheet_qty": sq, "invoice_qty": iq, "qty_diff": iq - sq, "group": r.get("group", ""),
             "master_kg": kg, "master_kgs": kgs, "invoice_kg": billed / pq if pq else None, "billed_kg": billed,
             "kg_source": source, "loaded_kg": None, "price_kg": None, "qty_kg": None,
             "master_price": kg * rate_per_kg if kg else None,
             "invoice_price": sum(l["qty"] * l["rate"] for l, _ in priced) / pq if pq else None}
        if kg is None:
            m["price_status"] = "NOT IN MASTER"
            m["qty_status"] = "NOT IN MASTER"
            not_in_master.append(m)
        else:
            m["loaded_kg"] = sq * kg
            m["qty_kg"] = (iq - sq) * kg
            m["price_kg"] = billed - iq * kg   # = sum over lines of qty x (invoice kg - its code's kg); keeps the split exact
            # unit price: each line's billed kg/pc (2 dp) against its own item code's master kg/pc; worst line decides
            gaps = [round(l["rate"] / rate_per_kg, 2) - round(k, 2) for l, k in priced if k]
            if source != "master":
                m["price_status"] = "NOT IN MASTER"
                m["note"] = (f"not in master - weighed at {kg:.2f} kg/pc from the "
                             + ("invoice rate" if source == "invoice" else f"invoice rate of {source[8:]} (similar model)"))
                not_in_master.append(m)
            elif not gaps:
                m["price_status"] = "NOT BILLED"
            else:
                worst = max(gaps, key=abs)
                m["kgpc_gap"] = round(worst, 2)
                m["price_status"] = "OK" if abs(worst) < 0.005 else ("PRICE HIGH" if worst > 0 else "PRICE LOW")
            m["qty_status"] = "OK" if iq == sq else ("MORE BILLED" if iq > sq else "LESS BILLED")
            if source == "master" and len(kgs) > 1:
                m["note"] = (f"item codes weigh {', '.join(f'{x:g}' for x in kgs)} kg - each invoice line uses its own; "
                             f"{kg:.2f} kg/pc average used for the sheet qty")
        models.append(m)

    known = [m for m in models if m["master_kg"] is not None]
    loaded = sum(m["loaded_kg"] for m in known)
    billed = sum(m["billed_kg"] for m in models)
    price = sum(m["price_kg"] for m in known)
    qty = sum(m["qty_kg"] for m in known)
    unknown_billed = sum(m["billed_kg"] for m in models if m["master_kg"] is None)
    # whole kg everywhere; the quantity part takes up any 1 kg rounding so the parts always add up on paper
    loaded_r, billed_r, price_r, nim_r = round(loaded), round(billed), round(price), round(unknown_billed)
    summary = {
        "rate_per_kg": rate_per_kg,
        "loaded_kg": loaded_r, "billed_kg": billed_r, "variance_kg": billed_r - loaded_r,
        "price_kg": price_r, "qty_kg": billed_r - loaded_r - price_r - nim_r, "not_in_master_kg": nim_r,
        "loaded_pcs": sum(m["sheet_qty"] for m in models), "billed_pcs": sum(m["invoice_qty"] for m in models),
        "loaded_pcs_weighed": sum(m["sheet_qty"] for m in known),   # sheet pcs with a master weight
        "price_issues": sum(m["price_status"] in ("PRICE HIGH", "PRICE LOW") for m in models),
        "qty_issues": sum(m["qty_status"] in ("MORE BILLED", "LESS BILLED") for m in models),
        "not_in_master": len(not_in_master),
        "sheet_not_in_master": [m for m in not_in_master if m["sheet_qty"]],   # weighed from the invoice, or not at all
    }
    # the parts must add up to the variance; anything else is a bug
    assert abs((price + qty + unknown_billed) - (billed - loaded)) < 0.01, "weight split does not add up"
    return {"summary": summary, "models": models}


# ---------------------------------------------------------------- Excel
def add_weight_sheets(wb, w, master):
    """Adds 'Weight reconciliation' (summary, Check 1, Check 2) and 'Master sheet (loaded)' tabs."""
    from openpyxl.styles import Font, PatternFill

    bold, title = Font(bold=True), Font(bold=True, size=14)
    red, green, amber = (PatternFill("solid", fgColor=c) for c in ("FFC7CE", "C6EFCE", "FFF2CC"))
    kgf = "#,##0"   # whole kg; kg/pc keeps 2 decimals
    s = w["summary"]

    ws = wb.create_sheet("Weight reconciliation", 0 if "Reconciliation" not in wb.sheetnames else 1)
    ws.append(["WEIGHT RECONCILIATION - loaded vs billed"])
    ws["A1"].font = title
    rows = [
        ("Today's Rate/Kg", s["rate_per_kg"]), ("Unit price match", "invoice kg/pc = master kg/pc (2 decimals)"),
        ("Loaded weight (sheet qty x master kg)", s["loaded_kg"]), ("Billed weight (invoice rate / Rate/Kg x qty)", s["billed_kg"]),
        ("VARIANCE (billed - loaded)", s["variance_kg"]),
        ("  of which Check 1 - unit price / weight", s["price_kg"]), ("  of which Check 2 - quantity", s["qty_kg"]),
        ("  of which invoice models not in master", s["not_in_master_kg"]),
    ]
    for label, v in rows:
        ws.append([label, v])
        ws.cell(ws.max_row, 1).font = bold
        if isinstance(v, (int, float)) and label != "Today's Rate/Kg":
            ws.cell(ws.max_row, 2).number_format = kgf
    ws.cell(6, 2).font = Font(bold=True, color="9C0006" if abs(s["variance_kg"]) >= 0.01 else "006100")
    ws.append(["Loaded pcs", s["loaded_pcs"], "Billed pcs", s["billed_pcs"]])
    ws.append([])

    def table(head, data, fmts, status_col=None):
        ws.append(head)
        for c in ws[ws.max_row]:
            c.font = bold
        for d in data:
            ws.append(d)
            for c, f in zip(ws[ws.max_row], fmts):
                if f:
                    c.number_format = f
            if status_col is not None:
                st = d[status_col]
                fill = green if st == "OK" else amber if st in ("NOT IN MASTER", "NOT BILLED") else red
                ws.cell(ws.max_row, status_col + 1).fill = fill

    models = sorted(w["models"], key=lambda m: model_sort_key(m["model"]))
    ws.append(["CHECK 1 - Unit price (run first): invoice rate vs master kg x Rate/Kg"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    table(["Model", "Status", "Master kg/pc", "Invoice kg/pc", "Master price", "Invoice price", "kg/pc gap", "Invoice qty", "Effect (kg)", "Note"],
          [[m["model"], m["price_status"], m["master_kg"], m["invoice_kg"], m["master_price"], m["invoice_price"],
            m.get("kgpc_gap"), m["invoice_qty"], m["price_kg"], m.get("note", "")] for m in models if m["invoice_qty"]],
          [None, None, "0.00", "0.00", kgf, kgf, "+0.00;-0.00;0.00", "0", kgf, None], status_col=1)
    ws.append([])
    ws.append(["CHECK 2 - Quantity: loading sheet vs invoice, converted to kg with master kg/pc"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
    table(["Model", "Status", "Sheet model", "Sheet line", "Sheet qty", "Invoice qty", "Diff (inv - sheet)", "Master kg/pc", "Effect (kg)", "Group"],
          [[m["model"], m["qty_status"], m["sheet_model"], m["sheet_lines"], m["sheet_qty"], m["invoice_qty"], m["qty_diff"],
            m["master_kg"], m["qty_kg"], m["group"]] for m in models],
          [None, None, None, None, "0", "0", "+0;-0;0", "0.00", kgf, None], status_col=1)
    for col, wd in zip("ABCDEFGHIJ", [44, 16, 14, 14, 14, 14, 16, 13, 14, 50]):
        ws.column_dimensions[col].width = wd

    # master sheet with column D = loaded qty from the loading sheet
    ms = wb.create_sheet("Master sheet (loaded)")
    ms.append(["Item code", "Material description", "weight (kg)", "Loaded qty (nos)", "Loaded kg", "Model"])
    for c in ms[1]:
        c.font = bold
    loaded = {m["model"]: m for m in w["models"] if m["kg_source"] == "master" and m["sheet_qty"]}
    done = set()
    for r in master["rows"]:
        m = loaded.get(r["model"])
        first = m is not None and r["model"] not in done   # a model's qty goes on its first master row
        if first:
            done.add(r["model"])
        ms.append([r["code"], r["desc"], r["kg"], m["sheet_qty"] if first else None,
                   round(m["loaded_kg"], 2) if first else None, r["model"]])   # same kg/pc as the reconciliation
        if first:
            ms.cell(ms.max_row, 4).fill = green
    # sheet models matched through an invoice item code but not by master description
    for key, m in loaded.items():
        if key not in done:
            ms.append(["(via invoice code)", key, m["master_kg"], m["sheet_qty"], round(m["loaded_kg"], 2), key])
            ms.cell(ms.max_row, 4).fill = green
    ms.append([])
    ms.append(["Loading sheet models NOT found in the master sheet (review)"])
    ms.cell(ms.max_row, 1).font = Font(bold=True, size=12)
    ms.append(["Model", "Written on sheet", "Sheet line", "Qty (nos)", "kg/pc used", "Loaded kg", "Weight from"])
    for c in ms[ms.max_row]:
        c.font = bold
    for m in sorted(w["summary"]["sheet_not_in_master"], key=lambda m: model_sort_key(m["model"])):
        ms.append([m["model"], m["sheet_model"], m["sheet_lines"], m["sheet_qty"], m["master_kg"],
                   round(m["loaded_kg"], 2) if m["loaded_kg"] is not None else None, m["kg_source"] or "no weight"])
        ms.cell(ms.max_row, 4).fill = amber
    for col, wd in zip("ABCDEF", [18, 46, 12, 16, 12, 14]):
        ms.column_dimensions[col].width = wd
    ms.freeze_panes = "A2"


def pdf_story(w, ss, small, colors, mm, Paragraph, Spacer, Table, TableStyle):
    """Weight reconciliation section for the combined PDF."""
    s = w["summary"]
    kg = lambda v: "" if v is None else f"{v:,.0f}"
    kgpc = lambda v: "" if v is None else f"{v:,.2f}"
    story = [Paragraph("Weight reconciliation · loaded vs billed", ss["Heading2"])]
    info = [["Loaded weight", kg(s["loaded_kg"]) + " kg", "Billed weight", kg(s["billed_kg"]) + " kg", "Variance", kg(s["variance_kg"]) + " kg"],
            ["Check 1 · unit price", kg(s["price_kg"]) + " kg", "Check 2 · quantity", kg(s["qty_kg"]) + " kg",
             "Not in master", kg(s["not_in_master_kg"]) + " kg"]]
    t = Table(info, hAlign="LEFT")
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                           *[("FONTNAME", (c, 0), (c, -1), "Helvetica-Bold") for c in (0, 2, 4)],
                           ("BACKGROUND", (5, 0), (5, 0), colors.HexColor("#FFC7CE" if abs(s["variance_kg"]) >= 0.01 else "#C6EFCE"))]))
    story += [t, Spacer(0, 4 * mm)]
    rows = [m for m in sorted(w["models"], key=lambda m: model_sort_key(m["model"]))
            if abs(m["price_kg"] or 0) >= 0.5 or abs(m["qty_kg"] or 0) >= 0.5 or m["master_kg"] is None]
    data = [["Model", "Sheet qty", "Invoice qty", "Master kg/pc", "Invoice kg/pc", "Check 1 kg", "Check 2 kg", "Status"]]
    for m in rows:
        data.append([m["model"], m["sheet_qty"], m["invoice_qty"], kgpc(m["master_kg"]), kgpc(m["invoice_kg"]),
                     kg(m["price_kg"]), kg(m["qty_kg"]), f"{m['price_status']} / {m['qty_status']}"])
    t = Table(data, repeatRows=1, hAlign="LEFT", colWidths=[w_ * mm for w_ in (34, 18, 20, 22, 22, 22, 22, 60)])
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (1, 1), (6, -1), "RIGHT"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7E9EE"))]))
    return story + [Paragraph("Models that contribute to the variance (similar models together):", small), t]
