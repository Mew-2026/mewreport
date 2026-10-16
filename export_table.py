"""Turn any table shown on the page into a formatted Excel or PDF file.

The page sends a spec: {"title", "subtitle", "columns": [{"name", "type"}], "rows": [[...]], "filename"}.
Column types: text, int (whole number), kg / money (whole number, thousands separator), kgpc (2 decimals),
status (coloured by its value). Numbers are written as numbers, so the Excel can be summed and filtered.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GOOD = {"CORRECT", "OK"}
MANUAL = {"MATCHED"}
PART = {"PARTLY MATCHED", "NOT IN MASTER", "NOT BILLED"}
FMT = {"int": "#,##0", "kg": "#,##0", "money": "#,##0", "kgpc": "0.00"}
FILL = {"good": "C6EFCE", "bad": "FFC7CE", "man": "DDEBF7", "part": "FFF2CC"}


def _tone(v):
    v = str(v or "").upper()
    if v in GOOD:
        return "good"
    if v in MANUAL:
        return "man"
    if v in PART:
        return "part"
    return "bad" if v else None


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def _validate(spec):
    cols = spec.get("columns") or []
    rows = spec.get("rows") or []
    if not cols or not isinstance(rows, list):
        raise ValueError("Nothing to export")
    if len(rows) > 5000:
        raise ValueError("Too many rows to export")
    rows = [list(r) + [None] * (len(cols) - len(r)) for r in rows]
    # leave out columns that are blank in every row (e.g. Group / Note when nothing was matched by hand)
    keep = [i for i in range(len(cols)) if any(r[i] not in (None, "") for r in rows)] or list(range(len(cols)))
    return [cols[i] for i in keep], [[r[i] for i in keep] for r in rows]


def to_xlsx(spec):
    cols, rows = _validate(spec)
    wb = Workbook()
    ws = wb.active
    ws.title = (spec.get("sheet") or "Table")[:31]
    ws.append([spec.get("title") or "Table"])
    ws["A1"].font = Font(bold=True, size=14)
    if spec.get("subtitle"):
        ws.append([spec["subtitle"]])
        ws.cell(ws.max_row, 1).font = Font(color="5A6A82")
    ws.append([f"Generated {datetime.now().strftime('%d-%m-%Y %H:%M')} · Mahesh Engineering Works"])
    ws.cell(ws.max_row, 1).font = Font(color="5A6A82", size=9)
    ws.append([])
    ws.append([c["name"] for c in cols])
    hdr = ws.max_row
    thin = Side(style="thin", color="CBD7EE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2373F4")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    for r in rows:
        out = []
        for c, v in zip(cols, list(r) + [None] * (len(cols) - len(r))):
            t = c.get("type", "text")
            n = _num(v) if t in FMT else None
            out.append(round(n) if (n is not None and t in ("int", "kg", "money")) else n if n is not None else ("" if v is None else v))
        ws.append(out)
        for c, cell in zip(cols, ws[ws.max_row]):
            t = c.get("type", "text")
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=t == "text", horizontal="right" if t in FMT else "left")
            if t in FMT:
                cell.number_format = FMT[t]
            if t == "status" and _tone(cell.value):
                cell.fill = PatternFill("solid", fgColor=FILL[_tone(cell.value)])
                cell.font = Font(bold=True)
    for i, c in enumerate(cols, start=1):
        longest = max([len(str(c["name"]))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[ws.cell(hdr, i).column_letter].width = max(8, min(48, longest + 3))
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    ws.auto_filter.ref = f"A{hdr}:{ws.cell(hdr, len(cols)).column_letter}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    cols, rows = _validate(spec)
    ss = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=ss["Normal"], fontSize=7.5, leading=9.5)
    cell_r = ParagraphStyle("cell_r", parent=cell, alignment=2)   # numbers right-aligned
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    page_w = landscape(A4)[0] - 20 * mm

    def fmt(v, t):
        n = _num(v) if t in FMT else None
        if n is None:
            return "" if v is None else str(v)
        return f"{n:,.2f}" if t == "kgpc" else f"{round(n):,}"

    data = [[Paragraph(c["name"], head) for c in cols]]
    style = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD7EE")),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2373F4")),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFF")])]
    for n, r in enumerate(rows, start=1):
        r = list(r) + [None] * (len(cols) - len(r))
        data.append([Paragraph(fmt(v, c.get("type", "text")).replace("&", "&amp;").replace("<", "&lt;"),
                               cell_r if c.get("type", "text") in FMT else cell) for c, v in zip(cols, r)])
        for i, c in enumerate(cols):
            if c.get("type") == "status" and _tone(r[i]):
                style.append(("BACKGROUND", (i, n), (i, n), colors.HexColor("#" + FILL[_tone(r[i])])))
    # column widths in proportion to their content, within the page
    weight = []
    for i, c in enumerate(cols):
        content = max([0] + [len(fmt(r[i], c.get("type", "text"))) for r in rows])
        word = max(len(w) for w in str(c["name"]).split())          # a header word never breaks mid-word
        status = 11 if c.get("type") == "status" else 0             # INCORRECT / PARTLY MATCHED fit on one line
        weight.append(max(word + 2, status, min(40, content + 1)))
    widths = [page_w * w / sum(weight) for w in weight]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                            bottomMargin=12 * mm, title=spec.get("title") or "Table", author="Mahesh Engineering Works")
    story = [Paragraph(spec.get("title") or "Table", ss["Title"])]
    if spec.get("subtitle"):
        story.append(Paragraph(spec["subtitle"].replace("&", "&amp;"), ss["Normal"]))
    story += [Spacer(0, 4 * mm)]
    t = Table(data, repeatRows=1, colWidths=widths)
    t.setStyle(TableStyle(style))
    story.append(t)

    def footer(canvas, doc_):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5A6A82"))
        canvas.drawString(10 * mm, 6 * mm, f"Mahesh Engineering Works · generated {datetime.now().strftime('%d-%m-%Y %H:%M')}")
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {doc_.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
