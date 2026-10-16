"""Match a handwritten loading sheet against the tax invoice(s) for the load and write an Excel report.

Usage:
  python reconcile.py --sheet loading_sheet.jpg --invoice inv1.pdf [inv2.pdf ...] [--out report.xlsx]
  python reconcile.py --sheet-json sheet.json --invoice inv1.pdf ...   # reuse an earlier OCR run (no credits)
"""
import argparse
import json
import os
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz

from invoice_parser import combine_invoices, model_sort_key, read_invoice

SUGGEST_SCORE = 75   # unmatched rows show the closest unmatched model above this as "possible match"


def normalize(model):
    """Sheet model -> the same key the invoice parser gives (invoice_parser.model_key).

    Only known handwriting / reading slips are corrected; nothing is guessed:
      ILTO / ILT0 -> ILTJ, ILT1 -> ILTT, UCTI -> UCTT, 4-digit Amaze ST -> STJ (J often lost),
      LPTT12150HL -> LPTT12150H (stray L), and a trailing PRO -> P (P at the end means Pro).
    """
    m = re.sub(r"[^A-Z0-9]", "", str(model).upper())
    m = re.sub(r"PRO$", "P", m)
    m = re.sub(r"^ILT[O0]", "ILTJ", m)
    m = re.sub(r"^ILT[1I]", "ILTT", m)
    m = re.sub(r"^UCT[I1]", "UCTT", m)
    m = re.sub(r"^(\d{4})ST(P?)$", r"\1STJ\2", m)   # only the 4-digit Amaze models (2042ST); RC18000ST stays
    m = re.sub(r"^(LPTT121[25]0H)L(P?)$", r"\1\2", m)
    return m


def suggest(key, candidates):
    """Closest model among the unmatched ones on the other side - shown as a hint, never applied."""
    best = max(((fuzz.ratio(key, c), c) for c in candidates), default=(0, None))
    return best[1] if best[0] >= SUGGEST_SCORE else ""


def refs(lines, multi):
    """Invoice Sr. Nos, grouped by invoice when the load has several invoices."""
    if not multi:
        return ", ".join(str(sr) for _, sr in lines)
    by_inv = defaultdict(list)
    for no, sr in lines:
        by_inv[no].append(str(sr))
    return "; ".join(f"{no}: {', '.join(srs)}" for no, srs in by_inv.items())


def reconcile(sheet, inv):
    """inv: result of invoice_parser.combine_invoices (one or more invoices for the load).

    Strict matching: a sheet model and an invoice model are the same battery only when their keys are
    equal after normalize(). Different numbers, a different family (RC / BC) or Pro vs non-Pro are never
    paired automatically; such rows stay NOT IN INVOICE / NOT IN SHEET with a "possible match" hint, and
    can be matched by hand on the page. Sheet rows with the same model are added up."""
    multi = len(inv["invoices"]) > 1
    agg = defaultdict(lambda: {"qty": 0, "lines": [], "codes": set(), "descs": set()})
    for l in inv["lines"]:
        a = agg[l["model"]]
        a["qty"] += l["qty"]
        a["lines"].append((l["invoice_no"], l["sr"]))
        a["codes"].add(l["code"])
        a["descs"].add(l["desc"])

    sheet_by_key = {}   # model key -> sheet rows with that model, in sheet order
    for r in sheet["rows"]:
        sheet_by_key.setdefault(normalize(r["model"]), []).append(r)
    unmatched_inv = [k for k in agg if k not in sheet_by_key]
    unmatched_sheet = [k for k in sheet_by_key if k not in agg]

    def row(key, sheet_rows):
        a = agg.get(key)
        sq = sum(r["qty"] for r in sheet_rows)
        iq = a["qty"] if a else 0
        if not sheet_rows:
            status, hint = "NOT IN SHEET", suggest(key, unmatched_sheet)
        elif not a:
            status, hint = "NOT IN INVOICE", suggest(key, unmatched_inv)
        else:
            status, hint = ("OK" if iq == sq else "QTY MISMATCH"), ""
        written = ", ".join(dict.fromkeys(r["model"] for r in sheet_rows))
        return {
            "line": sheet_rows[0]["line"] if sheet_rows else "",
            "sheet_lines": ", ".join(str(r["line"]) for r in sheet_rows),
            "sheet_model": written, "invoice_model": key if a else "", "model": key,
            "score": 100 if (a and sheet_rows) else "", "sheet_qty": sq, "invoice_qty": iq, "diff": iq - sq,
            "status": status, "result": "CORRECT" if status == "OK" else "INCORRECT",
            # blue: the sheet text was corrected to reach this model (e.g. ILTO18148 read as ILTJ18148)
            "check": bool(sheet_rows) and any(re.sub(r"[^A-Z0-9]", "", r["model"].upper()) != key for r in sheet_rows),
            "suggest": hint,
            "invoices": ", ".join(dict.fromkeys(no for no, _ in a["lines"])) if a else "",
            "invoice_lines": refs(a["lines"], multi) if a else "",
            "item_codes": ", ".join(sorted(a["codes"])) if a else "",
            "description": " / ".join(sorted(a["descs"])) if a else "",
        }

    out = [row(k, rows) for k, rows in sheet_by_key.items()]
    out += [row(k, []) for k in agg if k not in sheet_by_key]
    return out


def apply_groups(rows, groups):
    """Manual matches made on the page. groups: [{"rows": [row indexes], "note": str}].

    Rows grouped together are judged on their combined difference: net 0 -> MATCHED, otherwise
    PARTLY MATCHED with the remaining difference. Only rows that were INCORRECT can be grouped."""
    used = set()
    for n, g in enumerate(groups or [], 1):
        idx = sorted({int(i) for i in g.get("rows", [])})
        if len(idx) < 2:
            raise ValueError(f"Group {n}: choose at least 2 rows")
        for i in idx:
            if not 0 <= i < len(rows):
                raise ValueError(f"Group {n}: unknown row")
            if i in used:
                raise ValueError(f"Group {n}: a row can only be in one group")
            if rows[i]["result"] != "INCORRECT":
                raise ValueError(f"Group {n}: only incorrect rows can be matched by hand")
        used.update(idx)
        net = sum(rows[i]["diff"] for i in idx)
        note = " ".join(str(g.get("note") or "").split())[:200]
        for i in idx:
            rows[i].update(group=f"G{n}", group_net=net, note=note,
                           result="MATCHED" if net == 0 else "PARTLY MATCHED")
    return rows


RESULT_ORDER = {"INCORRECT": 0, "PARTLY MATCHED": 1, "MATCHED": 2, "CORRECT": 3}
RESULT_FILL = {"CORRECT": PatternFill("solid", fgColor="C6EFCE"), "INCORRECT": PatternFill("solid", fgColor="FFC7CE"),
               "MATCHED": PatternFill("solid", fgColor="DDEBF7"), "PARTLY MATCHED": PatternFill("solid", fgColor="FFF2CC")}
RESULT_FONT = {"CORRECT": Font(bold=True, color="006100"), "INCORRECT": Font(bold=True, color="9C0006"),
               "MATCHED": Font(bold=True, color="1F4E79"), "PARTLY MATCHED": Font(bold=True, color="7F6000")}
CHECK_FILL = PatternFill("solid", fgColor="BDD7EE")
BOLD = Font(bold=True)


def write_report(path, rows, sheet, inv):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"
    sheet_total = sum(r["qty"] for r in sheet["rows"])
    inv_total = sum(l["qty"] for l in inv["lines"])
    count = {k: sum(r["result"] == k for r in rows) for k in RESULT_ORDER}
    open_ = count["INCORRECT"] + count["PARTLY MATCHED"]

    ws.append(["LOADING SHEET vs INVOICE CHECK"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Overall", "ALL CORRECT" if not open_ else f"{open_} INCORRECT of {len(rows)}"])
    ws.append(["Loading sheet pcs", sheet_total, "Invoice pcs", inv_total, "Difference", inv_total - sheet_total])
    ws.append(["Correct rows", count["CORRECT"], "Incorrect rows", count["INCORRECT"],
               "Matched by hand", count["MATCHED"], "Partly matched", count["PARTLY MATCHED"]])
    for row in ws.iter_rows(min_row=2, max_row=4):
        for c in row[::2]:
            c.font = BOLD
    ws["B2"].font = RESULT_FONT["INCORRECT" if open_ else "CORRECT"]
    ws.append([])
    ws.append(["Invoice no", "Date", "Recipient", "Lines", "Pcs read", "Printed total", "Read check", "File"])
    for c in ws[ws.max_row]:
        c.font = BOLD
    for i in inv["invoices"]:
        read = sum(l["qty"] for l in i["lines"])
        printed = i.get("printed_total_qty")
        ok = printed is not None and printed == read
        ws.append([i["invoice_no"], i["date"], i["recipient"], len(i["lines"]), read, printed,
                   "OK" if ok else "CHECK: pcs read differ from printed total", i.get("file") or ""])
        ws.cell(ws.max_row, 7).font = RESULT_FONT["CORRECT" if ok else "INCORRECT"]
    for d in inv.get("duplicates", []):
        ws.append([f"Skipped duplicate upload: {d}"])
    ws.append([])

    cols = [("result", "Result"), ("status", "Reason"), ("sheet_lines", "Sheet line"),
            ("sheet_model", "Sheet model (as read)"), ("invoice_model", "Invoice model"), ("score", "Match %"),
            ("sheet_qty", "Sheet qty"), ("invoice_qty", "Invoice qty"), ("diff", "Diff (inv - sheet)"),
            ("group", "Group"), ("group_net", "Group net"), ("note", "Note"), ("suggest", "Possible match"),
            ("invoices", "Invoice no"), ("invoice_lines", "Invoice Sr. Nos"), ("item_codes", "Item codes"),
            ("description", "Invoice description")]
    ws.append([h for _, h in cols])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = BOLD
    # similar models next to each other (the Result column and the filter still show what needs attention)
    for r in sorted(rows, key=lambda r: model_sort_key(r.get("model") or r["invoice_model"] or r["sheet_model"])):
        ws.append([r.get(k, "") for k, _ in cols])
        cells = ws[ws.max_row]
        for c in cells[:9]:
            c.fill = RESULT_FILL[r["result"]]
        cells[0].font = RESULT_FONT[r["result"]]
        if r["check"]:
            for c in cells[3:6]:
                c.fill = CHECK_FILL
        for c in cells:
            c.alignment = Alignment(vertical="top")
    last = ws.max_row
    ws.append(["", "", "", "TOTAL", "", "", f"=SUM(G{hdr + 1}:G{last})", f"=SUM(H{hdr + 1}:H{last})",
               f"=SUM(I{hdr + 1}:I{last})"])
    for c in ws[ws.max_row]:
        c.font = BOLD
    ws.append([])
    ws.append(["Models are matched only when they are the same battery (P at the end = Pro). Blue model cells: the sheet text "
               "was corrected to reach the model (e.g. ILTO -> ILTJ). 'Possible match' is only a hint - match such rows by hand."])
    ws.append(["Compare 'Loading sheet pcs' with the total written on the paper: a misread row changes it."])
    groups = {}
    for r in rows:
        if r.get("group"):
            groups.setdefault(r["group"], []).append(r)
    if groups:
        ws.append([])
        ws.append(["Manual matches (made on the page)"])
        ws.cell(ws.max_row, 1).font = BOLD
        ws.append(["Group", "Result", "Net diff", "Rows (sheet model / invoice model : diff)", "", "", "", "", "", "Note"])
        for c in ws[ws.max_row]:
            c.font = BOLD
        for g, rs in sorted(groups.items(), key=lambda kv: int(kv[0][1:])):
            parts = ", ".join(f"{r['sheet_model'] or '-'} / {r['invoice_model'] or '-'} : {r['diff']:+d}" for r in rs)
            ws.append([g, rs[0]["result"], rs[0]["group_net"], parts, "", "", "", "", "", rs[0].get("note", "")])
            ws.cell(ws.max_row, 2).font = RESULT_FONT[rs[0]["result"]]
    ws.auto_filter.ref = f"A{hdr}:Q{last}"
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    for col, w in zip("ABCDEFGHIJKLMNOPQ", [16, 16, 10, 22, 16, 9, 10, 11, 15, 8, 10, 30, 16, 15, 22, 34, 60]):
        ws.column_dimensions[col].width = w

    il = wb.create_sheet("Invoice lines")
    il.append(["Invoice no", "Sr", "Item code", "Description", "Model", "Qty", "Rate", "Taxable value"])
    for l in inv["lines"]:
        il.append([l["invoice_no"], l["sr"], l["code"], l["desc"], l["model"], l["qty"], l["rate"], l["value"]])
    for c in il[1]:
        c.font = BOLD
    for col, w in zip("ABCDEFGH", [15, 6, 18, 48, 14, 7, 10, 14]):
        il.column_dimensions[col].width = w

    ls = wb.create_sheet("Loading sheet (OCR)")
    ls.append(["Line", "Row no. on sheet", "Model", "Qty"])
    for r in sheet["rows"]:
        ls.append([r["line"], r["sheet_row"], r["model"], r["qty"]])
    ls.append([])
    ls.append(["renamed.to replies per strip:"])
    for s in sheet.get("strips", []):
        ls.append([s["filename"]])
    for c in ls[1]:
        c.font = BOLD
    ls.column_dimensions["C"].width = 18
    wb.save(path)


def report_name(inv):
    nos = [i["invoice_no"] for i in inv["invoices"] if i["invoice_no"]]
    if not nos:
        return "reconciliation.xlsx"
    return "reconciliation_" + "_".join(nos[:3]) + ("_etc" if len(nos) > 3 else "") + ".xlsx"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--sheet", help="loading sheet: photo (read with renamed.to, uses credits) or .xlsx/.csv (Sr No | Model | Qty)")
    src.add_argument("--sheet-json", help="saved result of an earlier loading_sheet_ocr run")
    ap.add_argument("--invoice", required=True, nargs="+", help="one or more tax invoice PDFs for the load")
    ap.add_argument("--out", help="Excel report path (default: reconciliation_<invoice nos>.xlsx)")
    args = ap.parse_args()

    inv = combine_invoices([read_invoice(p) for p in args.invoice], [os.path.basename(p) for p in args.invoice])
    for d in inv["duplicates"]:
        print(f"skipped duplicate invoice: {d}")
    if args.sheet:
        if os.path.splitext(args.sheet)[1].lower() in (".xlsx", ".xlsm", ".csv"):
            from loading_sheet_excel import read_sheet_excel   # typed sheet: no OCR, no credits
            sheet = read_sheet_excel(args.sheet)
        else:
            from loading_sheet_ocr import read_sheet
            sheet = read_sheet(args.sheet)
            cache = os.path.splitext(args.sheet)[0] + ".ocr.json"
            with open(cache, "w", encoding="utf-8") as f:
                json.dump(sheet, f, indent=1)
            print(f"OCR saved to {cache} (rerun with --sheet-json to skip renamed.to)")
    else:
        with open(args.sheet_json, encoding="utf-8") as f:
            sheet = json.load(f)

    rows = reconcile(sheet, inv)
    out = args.out or report_name(inv)
    write_report(out, rows, sheet, inv)
    for r in rows:
        if r["result"] != "CORRECT" or r["check"]:
            flag = " (check model)" if r["check"] else ""
            print(f"{r['status']:15} {r['sheet_model'] or '-':14} -> {r['invoice_model'] or '-':12} "
                  f"sheet {r['sheet_qty']:>3} invoice {r['invoice_qty']:>3}{flag}")
    correct = sum(r["result"] == "CORRECT" for r in rows)
    print(f"{correct} correct, {len(rows) - correct} incorrect; sheet {sum(r['qty'] for r in sheet['rows'])} pcs, "
          f"invoices {sum(l['qty'] for l in inv['lines'])} pcs -> {out}")


if __name__ == "__main__":
    main()
