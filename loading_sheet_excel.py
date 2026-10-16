"""Read a loading sheet typed into Excel / CSV (Sr No | Model | Qty) - no OCR, no credits.

The header row is found by its labels (any row in the first 30, any column order). Rows end at
the first blank row after the data; rows without a quantity or with quantity 0 are skipped.
Returns the same shape as loading_sheet_ocr.read_sheet(), so the rest of the check is unchanged.

  python loading_sheet_excel.py --template public/loading_sheet_template.xlsx   # write the format file
"""
import csv
import io
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _header(row):
    cols = {}
    for i, v in enumerate(row):
        t = str(v or "").strip().lower()
        if not t:
            continue
        if ("model" in t or "battery" in t or "item" in t) and "model" not in cols:
            cols["model"] = i
        elif ("qty" in t or "quantity" in t or t in ("pcs", "nos")) and "qty" not in cols:
            cols["qty"] = i
        elif (t.startswith("sr") or t.startswith("s.no") or t in ("no", "no.", "#")) and "sr" not in cols:
            cols["sr"] = i
    return cols if "model" in cols and "qty" in cols else None


def _tables(src, name):
    if isinstance(src, (bytes, bytearray)):
        src = io.BytesIO(src)
    if os.path.splitext(name or "")[1].lower() == ".csv":
        data = src.read() if hasattr(src, "read") else open(src, "rb").read()
        yield "csv", list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
        return
    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except Exception:
        raise ValueError("Loading sheet: not an Excel file (.xlsx) or CSV")
    for ws in wb.worksheets:
        yield ws.title, [tuple(r) for r in ws.iter_rows(values_only=True)]


def read_sheet_excel(src, name=None):
    """src: path, bytes or file object; name: file name (decides CSV vs Excel)."""
    name = name or (src if isinstance(src, str) else "")
    for _, table in _tables(src, name):
        for h, row in enumerate(table[:30]):
            cols = _header(row)
            if not cols:
                continue
            rows, started = [], False
            for n, r in enumerate(table[h + 1:], start=h + 2):
                get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
                model = re.sub(r"\s+", "", str(get("model") or "")).upper()
                qty = _num(get("qty"))
                if not model and qty is None:
                    if started:
                        break  # blank row ends the list (notes may follow)
                    continue
                started = True
                if not model or not qty or qty <= 0:
                    continue
                if model in ("TOTAL", "GRANDTOTAL"):
                    break
                sr = _num(get("sr"))
                rows.append({"line": len(rows) + 1, "sheet_row": int(sr) if sr else n, "model": model,
                             "qty": int(round(qty))})
            if rows:
                return {"rows": rows, "total": sum(r["qty"] for r in rows), "strips": [], "source": "excel"}
            raise ValueError("Loading sheet: the Excel has the headers but no rows with a model and a quantity")
    raise ValueError("Loading sheet: no 'Model' and 'Qty' headers found - download the Excel format from the page")


EXAMPLE = [  # the 13/10/26 loading sheet, Patiala -> Muzaffarpur
    ("2042STJ", 3), ("2042TT", 2), ("2048STJ", 9), ("2048TT", 1), ("2060STJ", 6), ("2060TT", 17),
    ("2136TT", 1), ("2160TT", 5), ("2260TT", 3), ("4060TT", 1), ("5048TT", 12), ("5060TT", 1),
    ("ASTT12150H", 2), ("BC16048ST", 48), ("BC18048ST", 24), ("BC18060", 5), ("EC18060", 14),
    ("ILTJ18148", 18), ("ILTT18060", 3), ("ILTT18060P", 7), ("ILTT20060", 2), ("LM18075", 3),
    ("LPTT12120H", 34), ("LPTT12150H", 103), ("NEO18030", 1), ("PC18054P", 7), ("PC18054TJP", 8),
    ("RC18000", 1), ("RC18000ST", 7), ("RC18000STP", 3), ("RC24000P", 1), ("RC25000", 2),
    ("RC25000P", 14), ("SC16060", 40), ("SC18060", 10), ("TC13548ST", 8), ("TC15048ST", 5),
    ("TC15060", 1), ("TZ13548ST", 2), ("UCTT24066", 2), ("RC18048ST", 1), ("RC18000P", 2),
    ("ILTT24066", 1),
]


def write_template(path, rows=None):
    """The Excel format: headers plus rows. Without rows it is the downloadable format, filled with the
    13/10/26 sheet as a worked example; with rows (model, qty) it saves a checked / edited loading sheet.
    path may be a file path or a file object."""
    example = rows is None
    rows = EXAMPLE if example else rows
    wb = Workbook()
    ws = wb.active
    ws.title = "Loading sheet"
    thin = Side(style="thin", color="B7C4D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(["Sr No", "Model", "Qty"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F3D63")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for i, (m, q) in enumerate(rows, 1):
        ws.append([i, m, q])
        for c in ws[ws.max_row]:
            c.border = border
        ws.cell(ws.max_row, 1).alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 10
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    if not example:
        wb.save(path)
        return
    info = wb.create_sheet("How to fill")
    notes = [
        "Loading sheet - Excel format (Mahesh Engineering Works)",
        "",
        "1. Keep the header row: Sr No | Model | Qty.",
        "2. One battery model per row, the quantity loaded in Qty. Write the model as on the sheet,",
        "   e.g. 2042STJ, 2060TT, LPTT12150H, RC18000ST. Pro models end in P (RC24000P, ILTT18060P).",
        "3. The rows on the 'Loading sheet' tab are an example (13/10/26, Patiala -> Muzaffarpur, 440 pcs):",
        "   replace them with your own rows.",
        "4. Leave out crossed-out rows. Do not add a total row - the check adds the quantities itself.",
        "5. Save as .xlsx (or .csv) and choose it under Step 1 > Excel.",
    ]
    for t in notes:
        info.append([t])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 100
    wb.save(path)


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--template":
        write_template(sys.argv[2])
        print("wrote", sys.argv[2])
    elif len(sys.argv) == 2:
        s = read_sheet_excel(sys.argv[1])
        for r in s["rows"]:
            print(r["line"], r["model"], r["qty"], sep="\t")
        print("total", s["total"])
    else:
        sys.exit(__doc__)
