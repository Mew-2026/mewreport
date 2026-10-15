"""Check invoice rates against the mastersheet weight x today's rate per kg; write Excel and PDF.

Expected rate per piece = Rate/Kg x mastersheet weight (kg). Each invoice line's Rate(Rs) must be
within the tolerance (%) of it.

Usage:
  python rate_check.py --master mastersheet.xlsx --rate 119.54 --step1 reconciliation_x.xlsx   # combined report
  python rate_check.py --master mastersheet.xlsx --rate 119.54 --invoice inv1.pdf [inv2.pdf ...]
                       [--tolerance 0.5] [--out reports/rate_check]
"""
import argparse
import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz

from invoice_parser import combine_invoices, model_key, read_invoice
from mastersheet import code_key, read_master

DESC_SCORE = 90   # description fallback: minimum fuzzy score when the model number is not recognised


def _desc_norm(d):
    return " ".join(d.upper().replace("-", " ").split())


def _numbers(d):
    """Capacity / model numbers in a description (150, 2060, 18000 ...); version tags like V22 are ignored."""
    return set(re.findall(r"\d{3,}", d.replace(" ", "")))


def find_by_desc(master, desc):
    """Master row for an invoice line whose item code is not in the mastersheet. Descriptions of different
    batteries differ only in a few characters (2060TT / 2060STJ / 2160TT), so a row qualifies only when its
    battery model (Pro kept apart) and all its numbers are the same; the closest wording then wins.
    Returns (row, score) or (None, 0)."""
    key, d, nums = model_key(desc), _desc_norm(desc), _numbers(desc)
    known = key != desc.upper().replace(" ", "")   # model_key recognised a model
    pool = [m for m in master["rows"] if _numbers(m["desc"]) == nums and (m["model"] == key if known else True)]
    best = max(((fuzz.token_sort_ratio(d, _desc_norm(m["desc"])), m) for m in pool),
               key=lambda x: x[0], default=(0, None))
    floor = 0 if known else DESC_SCORE
    return (best[1], round(best[0])) if best[1] is not None and best[0] >= floor else (None, 0)


def check(master, inv, rate_per_kg, tol_pct):
    rows = []
    for l in inv["lines"]:
        m = master["by_code"].get(code_key(l["code"]))
        how, check_by_eye = "Item code", False
        if not m:
            m, score = find_by_desc(master, l["desc"])
            how, check_by_eye = (f"Description {score}%", True) if m else ("", False)
        rate, qty = l["rate"], l["qty"]
        r = {
            "invoice_no": l["invoice_no"], "sr": l["sr"], "code": l["code"], "desc": l["desc"], "qty": qty,
            "invoice_rate": rate, "invoice_value": l["value"], "matched_by": how, "check": check_by_eye,
            "master_code": m["code"] if m else "", "master_desc": m["desc"] if m else "",
            "kg": m["kg"] if m else None, "expected_rate": None, "diff": None, "diff_pct": None,
            "billed_per_kg": None, "expected_value": None, "value_diff": None, "total_kg": None,
        }
        if not m:
            r["status"] = "NOT IN MASTER"
        elif rate is None:
            r["status"] = "NO RATE ON INVOICE"
        else:
            exp = round(rate_per_kg * m["kg"], 2)
            diff = round(rate - exp, 2) + 0.0   # + 0.0 turns -0.0 into 0.0
            pct = diff / exp * 100 if exp else 0.0
            r.update(expected_rate=exp, diff=diff, diff_pct=round(pct, 2) + 0.0,
                     billed_per_kg=round(rate / m["kg"], 2), total_kg=round(m["kg"] * qty, 2),
                     expected_value=round(exp * qty, 2),
                     value_diff=round((l["value"] or rate * qty) - exp * qty, 2))
            r["status"] = "OK" if abs(pct) <= tol_pct + 1e-9 else ("RATE HIGH" if diff > 0 else "RATE LOW")
        r["result"] = "CORRECT" if r["status"] == "OK" else "INCORRECT"
        rows.append(r)
    return rows


def summary(rows, inv, rate_per_kg, tol_pct, master_name):
    correct = sum(r["result"] == "CORRECT" for r in rows)
    priced = [r for r in rows if r["expected_rate"] is not None]
    return {
        "rate_per_kg": rate_per_kg, "tolerance_pct": tol_pct, "master_file": master_name or "",
        "generated": datetime.now().strftime("%d-%m-%Y %H:%M"),
        "lines": len(rows), "correct": correct, "incorrect": len(rows) - correct,
        "not_in_master": sum(r["status"] == "NOT IN MASTER" for r in rows),
        "qty": sum(r["qty"] for r in rows),
        "total_kg": round(sum(r["total_kg"] for r in priced), 2),
        "invoice_value": round(sum(r["invoice_value"] or 0 for r in priced), 2),
        "expected_value": round(sum(r["expected_value"] for r in priced), 2),
        "value_diff": round(sum(r["value_diff"] for r in priced), 2),
        "invoices": [{k: i.get(k) for k in ("invoice_no", "date", "recipient", "file", "printed_total_qty")}
                     | {"lines": len(i["lines"]), "pcs": sum(l["qty"] for l in i["lines"])}
                     for i in inv["invoices"]],
        "duplicates": inv["duplicates"],
    }


def ordered(rows):
    """Incorrect lines first, then invoice order."""
    return sorted(rows, key=lambda r: (r["result"] == "CORRECT", r["invoice_no"] or "", r["sr"]))


# ---------------------------------------------------------------- Excel

RESULT_FILL = {"CORRECT": PatternFill("solid", fgColor="C6EFCE"), "INCORRECT": PatternFill("solid", fgColor="FFC7CE")}
RESULT_FONT = {"CORRECT": Font(bold=True, color="006100"), "INCORRECT": Font(bold=True, color="9C0006")}
CHECK_FILL = PatternFill("solid", fgColor="BDD7EE")
BOLD = Font(bold=True)
MONEY = "#,##0.00"

COLS = [  # key, header, width, number format
    ("result", "Result", 12, None), ("status", "Reason", 18, None), ("invoice_no", "Invoice no", 14, None),
    ("sr", "Sr", 5, None), ("code", "Item code", 16, None), ("desc", "Invoice description", 42, None),
    ("matched_by", "Matched by", 15, None), ("master_desc", "Master description", 42, None),
    ("kg", "Kg / pc", 9, "0.00"), ("qty", "Qty", 7, "0"), ("expected_rate", "Expected rate", 13, MONEY),
    ("invoice_rate", "Invoice rate", 13, MONEY), ("diff", "Diff (inv - exp)", 13, MONEY),
    ("diff_pct", "Diff %", 8, "0.00"), ("billed_per_kg", "Billed Rs/Kg", 11, MONEY),
    ("total_kg", "Total kg", 10, "#,##0.00"), ("expected_value", "Expected value", 15, MONEY),
    ("invoice_value", "Invoice taxable value", 16, MONEY), ("value_diff", "Value diff", 13, MONEY),
]


def write_excel(path, rows, s, master, wb=None):
    """wb: the Step 1 report workbook; the rate check is added to it as tab 2 (combined report)."""
    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Rate check"
    else:
        ws = wb.create_sheet("Rate check", 1)
    ws.append(["INVOICE RATE vs MASTERSHEET CHECK"])
    ws["A1"].font = Font(bold=True, size=14)
    all_ok = s["incorrect"] == 0
    ws.append(["Overall", "ALL CORRECT" if all_ok else f"{s['incorrect']} INCORRECT of {s['lines']}"])
    ws.append(["Today's Rate/Kg", s["rate_per_kg"], "Tolerance %", s["tolerance_pct"], "Mastersheet", s["master_file"]])
    ws.append(["Correct lines", s["correct"], "Incorrect lines", s["incorrect"], "Not in master", s["not_in_master"]])
    ws.append(["Total pcs", s["qty"], "Total kg", s["total_kg"], "Generated", s["generated"]])
    ws.append(["Expected value", s["expected_value"], "Invoice value", s["invoice_value"], "Difference", s["value_diff"]])
    for row in ws.iter_rows(min_row=2, max_row=6):
        for c in row[::2]:
            c.font = BOLD
    for ref in ("B6", "D6", "F6"):
        ws[ref].number_format = MONEY
    ws["B2"].font = RESULT_FONT["CORRECT" if all_ok else "INCORRECT"]
    ws.append([])

    ws.append(["Invoice no", "Date", "Recipient", "Lines", "Pcs read", "Printed total", "Read check", "File"])
    for c in ws[ws.max_row]:
        c.font = BOLD
    for i in s["invoices"]:
        ok = i["printed_total_qty"] is not None and i["printed_total_qty"] == i["pcs"]
        ws.append([i["invoice_no"], i["date"], i["recipient"], i["lines"], i["pcs"], i["printed_total_qty"],
                   "OK" if ok else "CHECK: pcs read differ from printed total", i["file"] or ""])
        ws.cell(ws.max_row, 7).font = RESULT_FONT["CORRECT" if ok else "INCORRECT"]
    for d in s["duplicates"]:
        ws.append([f"Skipped duplicate upload: {d}"])
    ws.append([])

    ws.append([h for _, h, _, _ in COLS])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in ordered(rows):
        ws.append([r[k] for k, _, _, _ in COLS])
        cells = ws[ws.max_row]
        for c, (_, _, _, fmt) in zip(cells, COLS):
            c.alignment = Alignment(vertical="top")
            if fmt:
                c.number_format = fmt
        for c in cells[:2]:
            c.fill = RESULT_FILL[r["result"]]
        cells[0].font = RESULT_FONT[r["result"]]
        for c in cells[10:14]:
            c.fill = RESULT_FILL[r["result"]]
        if r["check"]:
            for c in cells[6:8]:
                c.fill = CHECK_FILL
    last = ws.max_row
    col = {k: i + 1 for i, (k, _, _, _) in enumerate(COLS)}
    total = [""] * len(COLS)
    total[col["desc"] - 1] = "TOTAL"
    for k in ("qty", "total_kg", "expected_value", "invoice_value", "value_diff"):
        letter = ws.cell(hdr, col[k]).column_letter
        total[col[k] - 1] = f"=SUM({letter}{hdr + 1}:{letter}{last})"
    ws.append(total)
    for c, (_, _, _, fmt) in zip(ws[ws.max_row], COLS):
        c.font = BOLD
        if fmt:
            c.number_format = fmt
    ws.append([])
    ws.append([f"Expected rate = Today's Rate/Kg x mastersheet Kg/pc. CORRECT when the invoice rate is within "
               f"{s['tolerance_pct']}% of it."])
    ws.append(["Blue 'Matched by' cells: item code not in the mastersheet, matched by description - check by eye."])
    ws.auto_filter.ref = f"A{hdr}:{ws.cell(hdr, len(COLS)).column_letter}{last}"
    ws.freeze_panes = ws.cell(row=hdr + 1, column=3)
    for i, (_, _, w, _) in enumerate(COLS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    if master["problems"]:
        mp = wb.create_sheet("Mastersheet notes")
        mp.append(["Mastersheet notes: columns read from another position, rows skipped or duplicated"])
        mp["A1"].font = BOLD
        for p in master["problems"]:
            mp.append([p])
        mp.column_dimensions["A"].width = 110
    wb.save(path)


# ---------------------------------------------------------------- PDF

RECON_PDF_COLS = [  # Step 1 reconciliation columns repeated in the combined PDF: header, width (mm)
    ("Result", 18), ("Reason", 26), ("Sheet line", 14), ("Sheet model (as read)", 30), ("Invoice model", 28),
    ("Match %", 13), ("Sheet qty", 14), ("Invoice qty", 15), ("Diff (inv - sheet)", 23), ("Invoice Sr. Nos", 60),
]


def _recon_story(recon, ss, small, colors, mm, Paragraph, Spacer, Table, TableStyle):
    """Step 1 part of the combined PDF: its summary block and reconciliation table, as in the Step 1 Excel."""
    green, red = colors.HexColor("#C6EFCE"), colors.HexColor("#FFC7CE")
    story = [Paragraph("Step 1 · Loading sheet vs invoices", ss["Heading2"])]
    if recon["summary"]:
        width = max(len(r) for r in recon["summary"])
        info = [[str(v) for v in r] + [""] * (width - len(r)) for r in recon["summary"]]
        t = Table(info, hAlign="LEFT")
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                               *[("FONTNAME", (c, 0), (c, -1), "Helvetica-Bold") for c in range(0, width, 2)]]))
        story += [t, Spacer(0, 4 * mm)]
    idx = {h: i for i, h in enumerate(recon["head"])}
    cols = [(h, w) for h, w in RECON_PDF_COLS if h in idx]
    if not cols or not recon["rows"]:
        return story
    data = [[h for h, _ in cols]]
    style = [("FONTSIZE", (0, 0), (-1, -1), 6.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7E9EE"))]
    totals = {"Sheet qty": 0, "Invoice qty": 0, "Diff (inv - sheet)": 0}
    for n, r in enumerate(recon["rows"], start=1):
        cells = []
        for h, _ in cols:
            v = r[idx[h]]
            if h in totals and isinstance(v, (int, float)):
                totals[h] += v
            cells.append(Paragraph(str(v), small) if h == "Invoice Sr. Nos" and v else ("" if v is None else v))
        data.append(cells)
        if "Result" in idx:
            style.append(("BACKGROUND", (0, n), (1, n), green if r[idx["Result"]] == "CORRECT" else red))
    data.append([totals.get(h, "TOTAL" if h == "Sheet model (as read)" else "") for h, _ in cols])
    style.append(("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), "Helvetica-Bold"))
    t = Table(data, repeatRows=1, colWidths=[w * mm for _, w in cols], hAlign="LEFT")
    t.setStyle(TableStyle(style))
    return story + [t]


def write_pdf(path, rows, s, recon=None):
    """recon: Step 1 reconciliation (step1_report.read_step1); it goes first, the rate check on a new page."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=6.5, leading=8)
    green, red, blue = colors.HexColor("#C6EFCE"), colors.HexColor("#FFC7CE"), colors.HexColor("#BDD7EE")
    title = "Loading Sheet & Invoice Rate Check" if recon else "Invoice Rate vs Mastersheet Check"
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=title, author="mewreport")
    money = lambda v: "" if v is None else f"{v:,.2f}"
    story = [Paragraph(title, ss["Title"])]
    if recon:
        story += _recon_story(recon, ss, small, colors, mm, Paragraph, Spacer, Table, TableStyle)
        story += [PageBreak(), Paragraph("Step 2 · Invoice rate vs mastersheet", ss["Heading2"])]

    overall = "ALL CORRECT" if s["incorrect"] == 0 else f"{s['incorrect']} INCORRECT of {s['lines']} lines"
    info = [
        ["Overall", overall, "Today's Rate/Kg", money(s["rate_per_kg"]), "Tolerance", f"{s['tolerance_pct']}%"],
        ["Correct", s["correct"], "Incorrect", s["incorrect"], "Not in master", s["not_in_master"]],
        ["Total pcs", s["qty"], "Total kg", money(s["total_kg"]), "Generated", s["generated"]],
        ["Expected value", money(s["expected_value"]), "Invoice value", money(s["invoice_value"]),
         "Difference", money(s["value_diff"])],
    ]
    t = Table(info, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        *[("FONTNAME", (c, 0), (c, -1), "Helvetica-Bold") for c in (0, 2, 4)],
        ("BACKGROUND", (1, 0), (1, 0), green if s["incorrect"] == 0 else red),
    ]))
    story += [t, Spacer(0, 4 * mm)]

    inv = [["Invoice no", "Date", "Recipient", "Lines", "Pcs", "Printed total", "File"]]
    inv += [[i["invoice_no"], i["date"], Paragraph(i["recipient"] or "", small), i["lines"], i["pcs"],
             i["printed_total_qty"], Paragraph(i["file"] or "", small)] for i in s["invoices"]]
    t = Table(inv, hAlign="LEFT", colWidths=[28 * mm, 20 * mm, 70 * mm, 14 * mm, 14 * mm, 20 * mm, 80 * mm])
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                           ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(0, 4 * mm)]
    if s["duplicates"]:
        story.append(Paragraph("Skipped duplicate upload: " + ", ".join(s["duplicates"]), small))

    head = ["Result", "Reason", "Invoice / Sr", "Item code", "Invoice description", "Matched by", "Kg/pc",
            "Qty", "Expected rate", "Invoice rate", "Diff", "Diff %", "Value diff"]
    data = [head]
    style = [("FONTSIZE", (0, 0), (-1, -1), 6.5), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
             ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("ALIGN", (6, 1), (-1, -1), "RIGHT"), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7E9EE"))]
    for n, r in enumerate(ordered(rows), start=1):
        data.append([r["result"], r["status"], f"{r['invoice_no']} / {r['sr']}", r["code"],
                     Paragraph(r["desc"], small), Paragraph(r["matched_by"], small),
                     "" if r["kg"] is None else f"{r['kg']:.2f}", r["qty"], money(r["expected_rate"]),
                     money(r["invoice_rate"]), money(r["diff"]),
                     "" if r["diff_pct"] is None else f"{r['diff_pct']:.2f}", money(r["value_diff"])])
        fill = green if r["result"] == "CORRECT" else red
        style += [("BACKGROUND", (0, n), (1, n), fill), ("BACKGROUND", (8, n), (11, n), fill)]
        if r["check"]:
            style.append(("BACKGROUND", (5, n), (5, n), blue))
    data.append(["", "", "", "", "TOTAL", "", "", s["qty"], "", "", "", "", money(s["value_diff"])])
    style.append(("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), "Helvetica-Bold"))
    widths = [17, 24, 29, 23, 70, 20, 11, 9, 18, 18, 15, 11, 18]
    t = Table(data, repeatRows=1, colWidths=[w * mm for w in widths])
    t.setStyle(TableStyle(style))
    story += [t, Spacer(0, 3 * mm), Paragraph(
        f"Expected rate = Today's Rate/Kg x mastersheet Kg/pc. CORRECT when the invoice rate is within "
        f"{s['tolerance_pct']}% of it. Blue 'Matched by': item code not in the mastersheet, matched by "
        f"description - check by eye.", small)]

    def footer(canvas, doc_):
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(landscape(A4)[0] - 10 * mm, 6 * mm, f"Page {doc_.page}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- CLI

def report_name(inv, prefix="rate_check"):
    nos = [i["invoice_no"] for i in inv["invoices"] if i["invoice_no"]]
    if not nos:
        return prefix
    return prefix + "_" + "_".join(nos[:3]) + ("_etc" if len(nos) > 3 else "")


def _validate(rate_per_kg, tol_pct):
    if rate_per_kg <= 0:
        raise ValueError("Rate/Kg must be greater than 0")
    if tol_pct < 0:
        raise ValueError("Tolerance cannot be negative")


def run(master_src, master_name, invoices, names, rate_per_kg, tol_pct):
    """Rate check on parsed invoice PDFs. Returns (rows, summary, master, inv)."""
    _validate(rate_per_kg, tol_pct)
    master = read_master(master_src, master_name)
    inv = combine_invoices(invoices, names)
    rows = check(master, inv, rate_per_kg, tol_pct)
    return rows, summary(rows, inv, rate_per_kg, tol_pct, master_name), master, inv


def run_on_step1(step1_src, master_src, master_name, rate_per_kg, tol_pct):
    """Rate check on the invoice lines of a Step 1 report (reconciliation Excel).
    Returns (rows, summary, master, inv, step 1 workbook, step 1 reconciliation)."""
    from step1_report import read_step1
    _validate(rate_per_kg, tol_pct)
    wb, inv, recon = read_step1(step1_src)
    master = read_master(master_src, master_name)
    rows = check(master, inv, rate_per_kg, tol_pct)
    return rows, summary(rows, inv, rate_per_kg, tol_pct, master_name), master, inv, wb, recon


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--master", required=True, help="mastersheet .xlsx/.csv (Item code, Material description, weight (kg))")
    ap.add_argument("--rate", required=True, type=float, help="today's rate per kg (Rs)")
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--step1", help="Step 1 report (reconciliation .xlsx): writes the combined report")
    src.add_argument("--invoice", nargs="+", help="one or more tax invoice PDFs: writes the rate check only")
    ap.add_argument("--tolerance", type=float, default=0.5, help="allowed difference in %% (default 0.5)")
    ap.add_argument("--out", help="output path without extension; .xlsx and .pdf are written")
    args = ap.parse_args()

    if args.step1:
        rows, s, master, inv, wb, recon = run_on_step1(args.step1, args.master, os.path.basename(args.master),
                                                       args.rate, args.tolerance)
        out = args.out or report_name(inv, "report")
    else:
        rows, s, master, inv = run(args.master, os.path.basename(args.master),
                                   [read_invoice(p) for p in args.invoice],
                                   [os.path.basename(p) for p in args.invoice], args.rate, args.tolerance)
        wb = recon = None
        out = args.out or report_name(inv)
    if os.path.dirname(out):
        os.makedirs(os.path.dirname(out), exist_ok=True)
    write_excel(out + ".xlsx", rows, s, master, wb)
    write_pdf(out + ".pdf", rows, s, recon)

    for p in master["problems"]:
        print("mastersheet:", p)
    for d in s["duplicates"]:
        print("skipped duplicate invoice:", d)
    for r in ordered(rows):
        if r["result"] != "CORRECT" or r["check"]:
            flag = f" ({r['matched_by']} - check)" if r["check"] else ""
            exp = "-" if r["expected_rate"] is None else f"{r['expected_rate']:,.2f}"
            print(f"{r['status']:18} {r['invoice_no']}/{r['sr']:<3} {r['code']:15} expected {exp:>10} "
                  f"invoice {r['invoice_rate'] or 0:>10,.2f}{flag}")
    print(f"{s['correct']} correct, {s['incorrect']} incorrect of {s['lines']} lines -> {out}.xlsx, {out}.pdf")


if __name__ == "__main__":
    main()
