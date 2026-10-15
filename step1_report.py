"""Read a Step 1 report (reconciliation_*.xlsx written by reconcile.py) back in, so the rate check can run on
its invoice lines and the combined report can repeat its reconciliation table.
"""
import io

from openpyxl import load_workbook


def _s(v):
    """Invoice numbers may come back as numbers if the file was re-saved in Excel."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(v):
    try:
        return float(str(v).replace(",", "")) if v not in (None, "") else None
    except ValueError:
        return None


def read_step1(src):
    """src: path, bytes or file object. Returns (workbook, inv, recon).

    inv has the same shape as invoice_parser.combine_invoices(); recon holds the reconciliation summary
    and table for the PDF."""
    if isinstance(src, (bytes, bytearray)):
        src = io.BytesIO(src)
    try:
        wb = load_workbook(src)
    except Exception:
        raise ValueError("Step 1 report: not an Excel file")
    if "Invoice lines" not in wb.sheetnames or "Reconciliation" not in wb.sheetnames:
        raise ValueError("Step 1 report: this is not a Step 1 Excel (needs the 'Reconciliation' and "
                         "'Invoice lines' tabs)")
    for name in ("Rate check", "Mastersheet notes"):   # the file already went through step 2
        if name in wb.sheetnames:
            del wb[name]

    lines = []
    il = list(wb["Invoice lines"].iter_rows(values_only=True))
    head = [str(h or "").strip().lower() for h in il[0]] if il else []
    col = {k: head.index(k) for k in ("invoice no", "sr", "item code", "description", "model", "qty", "rate",
                                      "taxable value") if k in head}
    if len(col) < 8:
        raise ValueError("Step 1 report: the 'Invoice lines' tab is missing columns")
    for r in il[1:]:
        if r[col["qty"]] in (None, ""):
            continue
        lines.append({
            "invoice_no": _s(r[col["invoice no"]]), "sr": int(_num(r[col["sr"]]) or 0),
            "code": _s(r[col["item code"]]), "desc": _s(r[col["description"]]), "model": _s(r[col["model"]]),
            "qty": int(_num(r[col["qty"]]) or 0), "rate": _num(r[col["rate"]]), "value": _num(r[col["taxable value"]]),
        })
    if not lines:
        raise ValueError("Step 1 report: no invoice lines in it")

    rec = list(wb["Reconciliation"].iter_rows(values_only=True))
    summary, invoices, table_head, table = [], [], None, []
    i = 1
    while i < len(rec) and any(v not in (None, "") for v in rec[i]):   # summary block under the title
        summary.append([v for v in rec[i] if v not in (None, "")])
        i += 1
    for n, r in enumerate(rec):
        first = _s(r[0]) if r else ""
        if first == "Invoice no" and len(r) > 7 and _s(r[1]) == "Date":
            for x in rec[n + 1:]:
                if not x or x[0] in (None, "") or str(x[0]).startswith("Skipped"):
                    break
                invoices.append({"invoice_no": _s(x[0]), "date": _s(x[1]), "recipient": _s(x[2]) or None,
                                 "printed_total_qty": int(_num(x[5])) if _num(x[5]) is not None else None,
                                 "file": _s(x[7]) or None})
        if first == "Result" and table_head is None:
            table_head = [_s(v) for v in r if v not in (None, "")]
            for x in rec[n + 1:]:
                if not x or x[0] in (None, ""):
                    break
                table.append([x[k] for k in range(len(table_head))])
    known = {i["invoice_no"] for i in invoices}
    for no in dict.fromkeys(l["invoice_no"] for l in lines):   # invoice block missing: rebuild from lines
        if no not in known:
            invoices.append({"invoice_no": no, "date": "", "recipient": None, "printed_total_qty": None, "file": None})
    for inv in invoices:
        inv["lines"] = [l for l in lines if l["invoice_no"] == inv["invoice_no"]]

    inv = {"invoices": invoices, "lines": lines, "duplicates": []}
    recon = {"summary": summary, "head": table_head or [], "rows": table}
    return wb, inv, recon
